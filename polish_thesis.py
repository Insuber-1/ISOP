import sys,zipfile
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
root=Path(r'C:/project/project docx');master=root/'เล่มปริญญานิพนธ์_ฉบับรวม_ยังไม่เสร็จ.docx'
d=Document(str(master))
# Normalize every section to A4; use manual-specified margins by section role.
for i,s in enumerate(d.sections):
 s.page_width=Cm(21);s.page_height=Cm(29.7)
 if i==0: # hard cover layout: 2.54 cm all sides
  s.top_margin=Cm(2.54);s.left_margin=Cm(2.54);s.right_margin=Cm(2.54);s.bottom_margin=Cm(2.54)
 elif i==2: # approval page top margin in the manual
  s.top_margin=Cm(8.0);s.left_margin=Cm(3.81);s.right_margin=Cm(2.54);s.bottom_margin=Cm(2.54)
 else:
  s.top_margin=Cm(3.81);s.left_margin=Cm(3.81);s.right_margin=Cm(2.54);s.bottom_margin=Cm(2.54)
# Add 1.27cm before each chapter's title: text starts at the manual's 5.08cm on its first page,
# while subsequent pages keep the normal 3.81cm top margin.
section_index=0; first_after_break=True
for child in d._element.body.iterchildren():
 if child.tag==qn('w:p'):
  ppr=child.find(qn('w:pPr'))
  sect=ppr.find(qn('w:sectPr')) if ppr is not None else None
  if sect is not None:
   section_index+=1;first_after_break=True;continue
  if first_after_break:
   if 10 <= section_index <= 14:
    from docx.text.paragraph import Paragraph
    Paragraph(child,d._element.body).paragraph_format.space_before=Pt(36)
   first_after_break=False
# Set precise table widths for usable A4 text width ~14.65cm.
widths={0:[3.4,5.15,6.10],1:[1.15,2.15,4.15,4.25,2.95],2:[0.85,3.0,5.15,5.65],3:[1.2,8.3,5.15]}
for ti,t in enumerate(d.tables):
 ws=widths.get(ti)
 if not ws: continue
 t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 tblpr=t._tbl.tblPr
 layout=tblpr.find(qn('w:tblLayout'))
 if layout is None:layout=OxmlElement('w:tblLayout');tblpr.append(layout)
 layout.set(qn('w:type'),'fixed')
 tblw=tblpr.find(qn('w:tblW'))
 if tblw is None:tblw=OxmlElement('w:tblW');tblpr.append(tblw)
 tblw.set(qn('w:w'),str(round(sum(ws)*1440/2.54)));tblw.set(qn('w:type'),'dxa')
 for col,w in zip(t.columns,ws):col.width=Cm(w)
 for row_i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if row_i==0:
   repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trpr.append(repeat)
  for ci,cell in enumerate(row.cells):
   if ci<len(ws):cell.width=Cm(ws[ci])
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr()
   mar=tcpr.find(qn('w:tcMar'))
   if mar is None:mar=OxmlElement('w:tcMar');tcpr.append(mar)
   for side,val in [('top','55'),('bottom','55'),('start','90'),('end','90')]:
    e=mar.find(qn('w:'+side))
    if e is None:e=OxmlElement('w:'+side);mar.append(e)
    e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa')
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1
    if row_i==0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:
     run.font.name='TH SarabunPSK';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'TH SarabunPSK')
     if run.font.size is None:run.font.size=Pt(10)
     if row_i==0:run.bold=True
# Keep headings from being stranded at the bottom of a page.
for para in d.paragraphs:
 style=getattr(para.style,'name','') if para.style is not None else ''
 if style.startswith('Heading') or para.text.startswith(('บทที่ 1','บทที่ 2','บทที่ 3','บทที่ 4','บทที่ 5','เอกสารอ้างอิง','บรรณานุกรม','ภาคผนวก')):
  para.paragraph_format.keep_with_next=True
# Mark fields to refresh in Word, including the TOC.
upd=d.settings.element.find(qn('w:updateFields'))
if upd is None:upd=OxmlElement('w:updateFields');d.settings.element.append(upd)
upd.set(qn('w:val'),'true')
d.save(str(master))
# Apply compact A4 widths to standalone chapter drafts too; keep source documents untouched.
for path in [root/f'บทที่ {i}/ยังไม่เสร็จ_บทที่ {i}.docx' for i in range(2,6)]:
 x=Document(str(path))
 for sec in x.sections:sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(3.81);sec.left_margin=Cm(3.81);sec.right_margin=Cm(2.54);sec.bottom_margin=Cm(2.54)
 # chapter-opening heading begins 5.08 cm from paper edge; body pages use normal margins.
 if x.paragraphs:x.paragraphs[0].paragraph_format.space_before=Pt(36)
 for table in x.tables:
  count=len(table.columns)
  if count==3: ws=[3.4,5.15,6.1]
  elif count==5: ws=[1.15,2.15,4.15,4.25,2.95]
  else: continue
  table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
  for col,w in zip(table.columns,ws):col.width=Cm(w)
  for row in table.rows:
   for ci,cell in enumerate(row.cells):
    cell.width=Cm(ws[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for para in cell.paragraphs:
     para.paragraph_format.line_spacing=1;para.paragraph_format.space_after=Pt(0)
     for run in para.runs:
      if run.font.size is None:run.font.size=Pt(10)
      run.font.name='TH SarabunPSK';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'TH SarabunPSK')
 x.save(str(path))
# Validate all docs after saves and key layout dimensions.
for path in [master]+[root/f'บทที่ {i}/ยังไม่เสร็จ_บทที่ {i}.docx' for i in range(2,6)]:
 with zipfile.ZipFile(path) as z:
  assert z.testzip() is None and len(z.namelist())==len(set(z.namelist()))
  for n in z.namelist():
   if n.endswith(('.xml','.rels')):etree.fromstring(z.read(n))
 y=Document(str(path))
 assert all(abs(s.page_width-Cm(21))<100 and abs(s.page_height-Cm(29.7))<100 for s in y.sections)
print('polished',master,'bytes',master.stat().st_size,'sections',len(d.sections),'tables',len(d.tables),'images',len(d.inline_shapes))
print('all revised docs A4 and structurally valid')
